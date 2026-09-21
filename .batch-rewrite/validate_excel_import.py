# -*- coding: utf-8 -*-
"""
Dry-run validator that mirrors apps/server/src/services/excelImport.ts's row
validation WITHOUT touching the database. Used to prove the assembled
weekly-content-ja-kr-2026-09-21.xlsx will import with zero row errors before
actually calling POST /api/content/import-excel.
"""
import openpyxl
import os
import re

ROOT = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(ROOT, "weekly-content-ja-kr-2026-09-21.xlsx")

VALID_MARKETS = {"jp", "kr"}
VALID_SEGMENT_TYPES = {
    "morning_news", "nostalgia", "letter_reading", "quiz",
    "bedtime_story", "companionship", "gratitude_ritual",
}
IMPORTABLE_TIME_SLOTS = {"07:00", "12:00", "19:00", "21:00"}
REVIEW_REQUIRED_SEGMENT_TYPES = {"letter_reading", "companionship"}

WEEKLY_SCHEDULE_COLUMNS = [
    "market", "scheduled_date", "time_slot", "segment_type", "title", "script_text",
    "tags", "needs_review", "shorts_snippet", "target_keyword", "seo_title",
    "seo_description", "seo_tags", "hashtags",
]
QUIZ_QUESTIONS_COLUMNS = [
    "market", "date", "question", "choice_1", "choice_2", "choice_3", "choice_4",
    "answer", "fun_fact",
]


def validate_seo_title(market, seo_title, target_keyword):
    if not target_keyword or not seo_title:
        return True
    if market == "kr":
        first_five = " ".join(seo_title.strip().split()[:5])
        return target_keyword in first_five
    return target_keyword in seo_title[:20]


def read_sheet(ws, columns):
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cells = dict(zip(columns, row))
        if any(v not in (None, "") for v in cells.values()):
            rows.append((i, cells))
    return rows


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    assert "Weekly_Schedule" in wb.sheetnames, "Weekly_Schedule sheet missing"
    sched_ws = wb["Weekly_Schedule"]
    quiz_ws = wb["Quiz_Questions"] if "Quiz_Questions" in wb.sheetnames else None

    header = [c.value for c in sched_ws[1]]
    assert header == WEEKLY_SCHEDULE_COLUMNS, f"Weekly_Schedule header mismatch: {header}"
    if quiz_ws:
        qheader = [c.value for c in quiz_ws[1]]
        assert qheader == QUIZ_QUESTIONS_COLUMNS, f"Quiz_Questions header mismatch: {qheader}"

    sched_rows = read_sheet(sched_ws, WEEKLY_SCHEDULE_COLUMNS)
    quiz_rows = read_sheet(quiz_ws, QUIZ_QUESTIONS_COLUMNS) if quiz_ws else []

    errors = []
    seo_seen = {}
    quiz_item_count = 0
    for row_no, c in sched_rows:
        row_errs = []
        market = str(c["market"] or "").strip().lower()
        segment_type = str(c["segment_type"] or "").strip().lower()
        time_slot = str(c["time_slot"] or "").strip()
        title = str(c["title"] or "").strip()
        script_text = str(c["script_text"] or "").strip()
        date = str(c["scheduled_date"] or "").strip()

        if market not in VALID_MARKETS:
            row_errs.append(f"market invalid: {c['market']!r}")
        if not re.match(r"^\d{4}-\d{2}-\d{2}", date):
            row_errs.append(f"scheduled_date invalid: {c['scheduled_date']!r}")
        if time_slot not in IMPORTABLE_TIME_SLOTS:
            row_errs.append(f"time_slot invalid: {time_slot!r}")
        if segment_type not in VALID_SEGMENT_TYPES:
            row_errs.append(f"segment_type invalid: {segment_type!r}")
        if not title:
            row_errs.append("title required")
        if not script_text:
            row_errs.append("script_text required")

        target_keyword = str(c["target_keyword"] or "").strip()
        seo_title = str(c["seo_title"] or "").strip()
        if seo_title and not validate_seo_title(market, seo_title, target_keyword):
            row_errs.append(f"seo_title must contain target_keyword near the front: {seo_title!r} / kw={target_keyword!r}")

        seo_desc = str(c["seo_description"] or "").strip()
        if seo_desc:
            seo_seen.setdefault(seo_desc, []).append((market, date, segment_type, title))

        if segment_type == "quiz":
            quiz_item_count += 1

        if row_errs:
            errors.append((row_no, title, row_errs))

    dup_desc = {k: v for k, v in seo_seen.items() if len(v) > 1}

    # cross-check quiz linkage: every quiz Weekly_Schedule row should have >=1 matching Quiz_Questions rows
    quiz_keys_present = set()
    for _, c in quiz_rows:
        quiz_keys_present.add((str(c["market"]).strip().lower(), str(c["date"]).strip()[:10]))
    quiz_link_errors = []
    for row_no, c in sched_rows:
        if str(c["segment_type"] or "").strip().lower() != "quiz":
            continue
        key = (str(c["market"]).strip().lower(), str(c["scheduled_date"]).strip()[:10])
        if key not in quiz_keys_present:
            quiz_link_errors.append((row_no, c["title"], key))

    print(f"Weekly_Schedule rows read: {len(sched_rows)}")
    print(f"Quiz_Questions rows read: {len(quiz_rows)}")
    print(f"Quiz Weekly_Schedule items: {quiz_item_count}")
    print(f"Row validation errors: {len(errors)}")
    for row_no, title, errs in errors:
        print(f"  row {row_no} ({title}): {errs}")
    print(f"Duplicate seo_description across different videos: {len(dup_desc)}")
    for desc, owners in dup_desc.items():
        print(f"  {desc[:50]!r} used by {owners}")
    print(f"Quiz rows missing linked Quiz_Questions: {len(quiz_link_errors)}")
    for row_no, title, key in quiz_link_errors:
        print(f"  row {row_no} ({title}) key={key}")

    ok = not errors and not quiz_link_errors
    print("\nRESULT:", "IMPORT-READY (0 row errors)" if ok else "BLOCKED — fix errors above before importing")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
